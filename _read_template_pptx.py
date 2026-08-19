import os
from pptx import Presentation

template_path = r"C:\Users\pokus\Downloads\Telegram Desktop\Документы для Симулятора\Документы для Симулятора\Макет презентации.pptx"

prs = Presentation(template_path)

print(f"Slide dimensions: {prs.slide_width.inches:.3f} x {prs.slide_height.inches:.3f} inches")
print(f"Total slides: {len(prs.slides)}\n")

for i, slide in enumerate(prs.slides):
    print(f"==================================================")
    print(f"SLIDE {i+1}")
    print(f"==================================================")
    
    # Iterate through all shapes
    for shape_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                txt = p.text.strip()
                if txt:
                    # Replace line breaks and clean up
                    cleaned = " ".join(txt.split())
                    print(f"  [Shape {shape_idx}]: {cleaned}")
                    
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            print(f"  [Notes]: {notes}")
    print()
